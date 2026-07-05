"""Lees een Excel-model uit en dump de volledige structuur naar JSON/Markdown.

Doel: TSG Model.xlsx (en Yield Report Database.xlsx) machinaal uitlezen
zodat we de exacte formules, named ranges en aannames kunnen overzetten
naar engine/. Gebruik:

    python tools/extract_excel.py "data/TSG Model.xlsx"

Output komt naast het bronbestand te staan:
  - <naam>.structuur.json  — alle cellen met waarde én formule, per werkblad
  - <naam>.overzicht.md    — leesbaar overzicht (werkbladen, named ranges, formules)
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl


def extraheer(pad: Path) -> dict:
    # data_only=False geeft de formules; een tweede load geeft de laatst
    # berekende waarden (zoals Excel ze opsloeg) voor golden tests.
    wb_formules = openpyxl.load_workbook(pad, data_only=False)
    wb_waarden = openpyxl.load_workbook(pad, data_only=True)

    resultaat: dict = {"bestand": pad.name, "werkbladen": {}, "named_ranges": {}}

    for naam in wb_formules.defined_names:
        dn = wb_formules.defined_names[naam]
        resultaat["named_ranges"][naam] = dn.attr_text

    for ws_f in wb_formules.worksheets:
        ws_w = wb_waarden[ws_f.title]
        cellen = {}
        for rij in ws_f.iter_rows():
            for cel in rij:
                if cel.value is None:
                    continue
                info: dict = {}
                if isinstance(cel.value, str) and cel.value.startswith("="):
                    info["formule"] = cel.value
                    info["waarde"] = ws_w[cel.coordinate].value
                else:
                    info["waarde"] = cel.value
                cellen[cel.coordinate] = info
        resultaat["werkbladen"][ws_f.title] = cellen

    return resultaat


def schrijf_overzicht(data: dict, pad: Path) -> None:
    regels = [f"# Structuur van {data['bestand']}", ""]
    if data["named_ranges"]:
        regels += ["## Named ranges", ""]
        for naam, ref in sorted(data["named_ranges"].items()):
            regels.append(f"- `{naam}` -> `{ref}`")
        regels.append("")
    for blad, cellen in data["werkbladen"].items():
        formules = {c: i for c, i in cellen.items() if "formule" in i}
        regels += [f"## Werkblad: {blad}", "", f"{len(cellen)} gevulde cellen, {len(formules)} formules", ""]
        for coord, info in formules.items():
            regels.append(f"- `{coord}`: `{info['formule']}` = `{info['waarde']}`")
        regels.append("")
    pad.write_text("\n".join(regels), encoding="utf-8")


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit("Gebruik: python tools/extract_excel.py <pad-naar-xlsx>")
    bron = Path(sys.argv[1])
    data = extraheer(bron)

    json_pad = bron.with_suffix(".structuur.json")
    json_pad.write_text(json.dumps(data, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    md_pad = bron.with_suffix(".overzicht.md")
    schrijf_overzicht(data, md_pad)
    print(f"Geschreven: {json_pad}\nGeschreven: {md_pad}")


if __name__ == "__main__":
    main()
